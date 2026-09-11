import io
import json
import sqlite3
from pathlib import Path
from datetime import datetime
from collections import Counter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

STATES=['已交','未交','待复核','请假','缺勤','免交']

def connect(path):
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    db=sqlite3.connect(path)
    db.row_factory=sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    db.executescript('''
    CREATE TABLE IF NOT EXISTS students(class TEXT, sid TEXT, name TEXT NOT NULL, pocket TEXT,source_class TEXT NOT NULL DEFAULT '',
      PRIMARY KEY(class,sid), UNIQUE(class,pocket));
    CREATE TABLE IF NOT EXISTS sessions(id INTEGER PRIMARY KEY, class TEXT, term TEXT, day TEXT, lesson TEXT,
      photo TEXT, calibration TEXT, UNIQUE(class,term,day,lesson));
    CREATE TABLE IF NOT EXISTS records(session INTEGER REFERENCES sessions(id),sid TEXT,name TEXT,pocket TEXT,
      prediction TEXT,score REAL,state TEXT,note TEXT,PRIMARY KEY(session,sid));
    CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, time TEXT, session INTEGER, sid TEXT, before TEXT, after TEXT);
    CREATE TABLE IF NOT EXISTS seeds(filename TEXT PRIMARY KEY);
    ''')
    if 'source_class' not in {r['name'] for r in db.execute('PRAGMA table_info(students)')}:
        # Migrate the initial prototype without changing any saved lesson records.
        with db:
            db.execute('ALTER TABLE students RENAME TO students_v0')
            db.execute("CREATE TABLE students(class TEXT,sid TEXT,name TEXT NOT NULL,pocket TEXT,source_class TEXT NOT NULL DEFAULT '',PRIMARY KEY(class,sid),UNIQUE(class,pocket))")
            db.execute('INSERT INTO students SELECT class,sid,name,pocket,class FROM students_v0')
            db.execute('DROP TABLE students_v0')
    return db

def seed_rosters(db,paths):
    for path in paths:
        marker='current-class-v2:'+path.name
        if db.execute('SELECT 1 FROM seeds WHERE filename=?',(marker,)).fetchone(): continue
        incoming=read_roster(path)
        # The user confirmed that the provided files represent current classes;
        # original Excel class values for transferred students are retained.
        target=Counter(r['class_name'] for r in incoming).most_common(1)[0][0]
        incoming=assign_current_class(incoming,target)
        import_roster(db,incoming)
        with db:
            for r in incoming:
                if r['source_class']!=target:
                    db.execute('DELETE FROM students WHERE class=? AND sid=?',(r['source_class'],r['sid']))
            db.execute('INSERT INTO seeds VALUES(?)',(marker,))

def assign_current_class(rows,target):
    """Keep existing native assignments; conflicting transfers remain unassigned."""
    target=target.strip()
    if not target: raise ValueError('当前班级名称不能为空。')
    occupied=set(); result=[]
    for r in sorted(rows,key=lambda r:r['class_name']!=target):
        p=r['pocket']
        if p in occupied: p=None
        if p: occupied.add(p)
        result.append(dict(r,class_name=target,source_class=r.get('source_class',r['class_name']),pocket=p))
    validate_students(result)
    return result

def read_roster(source):
    wb=load_workbook(io.BytesIO(source) if isinstance(source,bytes) else source,read_only=True,data_only=True)
    result=[]
    try:
        for ws in wb:
            rows=iter(ws.values)
            header=next(rows,())
            names=[str(v).strip() if v is not None else '' for v in header]
            if not all(k in names for k in ['学号/工号','姓名','班级']):
                continue
            for n,row in enumerate(rows,2):
                if not any(v is not None for v in row): continue
                vals=[row[names.index(k)] for k in ['学号/工号','姓名','班级']]
                if any(v is None or not str(v).strip() for v in vals):
                    raise ValueError(f'{ws.title} 第 {n} 行缺少学号、姓名或班级。')
                sid,name,cls=[str(v).strip() for v in vals]
                if not sid.isdigit() or len(sid)<2:
                    raise ValueError(f'第 {n} 行学号须为至少两位数字文本。')
                pocket=sid[-2:]
                if '袋号' in names:
                    value=row[names.index('袋号')]
                    pocket=None if value is None or str(value).strip() in ('','待分配') else f'{int(value):02}'
                if pocket is not None and not 1<=int(pocket)<=54:
                    raise ValueError(f'{cls} / {sid} 的尾号不在 01～54，请在导入表中增加有效“袋号”列。')
                result.append(dict(class_name=cls,sid=sid,name=name,pocket=pocket,source_class=cls))
    finally:
        wb.close()
    validate_students(result)
    return result

def validate_students(rows):
    if not rows: raise ValueError('未发现学生；首行须包含学号/工号、姓名、班级。')
    for r in rows:
        if not r['class_name'].strip() or not r['name'].strip() or not r['sid'].strip(): raise ValueError('班级、姓名、学号不能为空。')
        if r['pocket'] is not None and (len(r['pocket'])!=2 or not r['pocket'].isdigit() or not 1<=int(r['pocket'])<=54): raise ValueError('袋号必须为 01～54，或留空待分配。')
    for key,label in [('sid','学号'),('pocket','袋号')]:
        count=Counter((r['class_name'],r[key]) for r in rows if r[key] is not None)
        duplicate=next((k for k,v in count.items() if v>1),None)
        if duplicate: raise ValueError(f'{duplicate[0]} 内{label}重复：{duplicate[1]}。')

def import_roster(db,rows):
    """Merge named students; keep students absent from the file, validate final class first."""
    validate_students(rows)
    classes=sorted({r['class_name'] for r in rows})
    merged=[]
    for cls in classes:
        existing={r['sid']:dict(class_name=cls,sid=r['sid'],name=r['name'],pocket=r['pocket'],source_class=r['source_class']) for r in students(db,cls)}
        existing.update({r['sid']:r for r in rows if r['class_name']==cls})
        merged.extend(existing.values())
    validate_students(merged)
    with db:
        for cls in classes: db.execute('DELETE FROM students WHERE class=?',(cls,))
        db.executemany('INSERT INTO students VALUES(?,?,?,?,?)',[(r['class_name'],r['sid'],r['name'],r['pocket'],r.get('source_class',r['class_name'])) for r in merged])
    return classes

def students(db,cls):
    return [dict(r) for r in db.execute('SELECT * FROM students WHERE class=? ORDER BY pocket IS NULL,pocket,sid',(cls,))]

def classes(db):
    return [r[0] for r in db.execute('SELECT DISTINCT class FROM students ORDER BY class')]

def save_session(db,cls,term,day,lesson,photo,calibration,records):
    if not term.strip() or not lesson.strip(): raise ValueError('学期和课次不能为空。')
    if len({r['sid'] for r in records})!=len(records): raise ValueError('同一课次学号重复。')
    if any(r['state'] not in STATES for r in records): raise ValueError('状态无效。')
    with db:
        db.execute('INSERT INTO sessions(class,term,day,lesson,photo,calibration) VALUES(?,?,?,?,?,?) ON CONFLICT(class,term,day,lesson) DO UPDATE SET photo=excluded.photo,calibration=excluded.calibration',
                   (cls,term.strip(),day,lesson.strip(),photo,json.dumps(calibration,ensure_ascii=False)))
        session=db.execute('SELECT id FROM sessions WHERE class=? AND term=? AND day=? AND lesson=?',(cls,term.strip(),day,lesson.strip())).fetchone()[0]
        for r in records:
            old=db.execute('SELECT * FROM records WHERE session=? AND sid=?',(session,r['sid'])).fetchone()
            # Existing records keep the historical identity/pocket snapshot.
            name,pocket=(old['name'],old['pocket']) if old else (r['name'],r['pocket'])
            after=dict(r,name=name,pocket=pocket)
            db.execute('INSERT INTO audit(time,session,sid,before,after) VALUES(?,?,?,?,?)',(datetime.now().isoformat(),session,r['sid'],json.dumps(dict(old),ensure_ascii=False) if old else None,json.dumps(after,ensure_ascii=False)))
            db.execute('INSERT INTO records VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(session,sid) DO UPDATE SET prediction=excluded.prediction,score=excluded.score,state=excluded.state,note=excluded.note',
                       (session,r['sid'],name,pocket,r['prediction'],r.get('score',0),r['state'],r.get('note','')))
    return session

def correct_record(db,session,sid,state,note):
    if state not in STATES: raise ValueError('状态无效。')
    old=db.execute('SELECT * FROM records WHERE session=? AND sid=?',(session,sid)).fetchone()
    if not old: raise ValueError('记录不存在。')
    after=dict(old,state=state,note=note)
    with db:
        db.execute('UPDATE records SET state=?,note=? WHERE session=? AND sid=?',(state,note,session,sid))
        db.execute('INSERT INTO audit(time,session,sid,before,after) VALUES(?,?,?,?,?)',(datetime.now().isoformat(),session,sid,json.dumps(dict(old),ensure_ascii=False),json.dumps(after,ensure_ascii=False)))

def details(db,cls=None,term=None,start=None,end=None):
    sql='SELECT s.id AS session,s.class,s.term,s.day,s.lesson,s.photo,r.* FROM records r JOIN sessions s ON s.id=r.session WHERE 1=1'
    args=[]
    for key,op,value in [('s.class','=',cls),('s.term','=',term),('s.day','>=',start),('s.day','<=',end)]:
        if value: sql+=f' AND {key}{op}?'; args.append(value)
    return [dict(r) for r in db.execute(sql+' ORDER BY s.day,s.lesson,r.pocket',args)]

def summary(db,rows,cls=None):
    result={}
    roster=students(db,cls) if cls else [dict(r) for r in db.execute('SELECT * FROM students')]
    for r in roster:
        result[(r['class'],r['sid'])]={'班级':r['class'],'学号':r['sid'],'姓名':r['name'],**{s+'次数':0 for s in STATES}}
    for r in rows:
        key=(r['class'],r['sid'])
        if key not in result: result[key]={'班级':r['class'],'学号':r['sid'],'姓名':r['name'],**{s+'次数':0 for s in STATES}}
        result[key][r['state']+'次数']+=1
    return sorted(result.values(),key=lambda r:(r['班级'],-r['未交次数'],r['学号']))

def export_excel(totals,rows,filters):
    wb=Workbook(); wb.remove(wb.active)
    detail=[{'班级':r['class'],'学期':r['term'],'日期':r['day'],'课次':r['lesson'],'袋号':r['pocket'],'学号':r['sid'],'姓名':r['name'],'最终状态':r['state'],'备注':r['note']} for r in rows]
    for title,data,headers in [('汇总',totals,['班级','学号','姓名']+[s+'次数' for s in STATES]),('明细',detail,['班级','学期','日期','课次','袋号','学号','姓名','最终状态','备注'])]:
        ws=wb.create_sheet(title); ws.append([filters]); ws.append(headers)
        for item in data:
            ws.append([item.get(k,'') for k in headers])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str): cell.data_type='s'
        for cell in ws[2]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='087F8C')
        ws.freeze_panes='D3'; ws.auto_filter.ref=f'A2:{ws.cell(max(2,ws.max_row),len(headers)).coordinate}'
        for i,k in enumerate(headers,1): ws.column_dimensions[ws.cell(2,i).column_letter].width=38 if k=='班级' else (26 if k=='备注' else 18)
        ws.row_dimensions[1].height=28
        for row in ws.iter_rows(min_row=3):
            for cell in row: cell.alignment=Alignment(vertical='center',wrap_text=True)
    out=io.BytesIO(); wb.save(out); return out.getvalue()
