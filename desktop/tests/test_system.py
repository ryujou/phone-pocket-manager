import io
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pytest
from openpyxl import load_workbook,Workbook
import storage as s
from vision import read_image,recognize,rectify
ROOT=Path(__file__).resolve().parents[1]

@pytest.fixture
def db(tmp_path):
    db=s.connect(tmp_path/'db.sqlite'); yield db; db.close()

def roster(cls='A',sid='90000101',name='测试甲',pocket='01'):
    return dict(class_name=cls,sid=sid,name=name,pocket=pocket)

def record(state='未交'):
    return dict(sid='90000101',name='测试甲',pocket='01',prediction='empty',score=.9,state=state,note='')

def test_mock_roster_seed_once(db):
    paths=sorted((ROOT/'samples').glob('roster*.xlsx'))
    s.seed_rosters(db,paths)
    assert len(s.classes(db))==2
    assert db.execute('SELECT COUNT(*) FROM students').fetchone()[0]==12
    assert s.students(db,'示例一班')[0]['name']=='张三'
    s.seed_rosters(db,paths)
    assert db.execute('SELECT COUNT(*) FROM students').fetchone()[0]==12
    assert s.details(db)==[]

def test_multiclass_import_and_rollback(db):
    s.import_roster(db,[roster(),roster('B')]); assert len(s.classes(db))==2
    s.import_roster(db,[roster(name='新名字')]); assert s.students(db,'B')[0]['name']=='测试甲'
    with pytest.raises(ValueError): s.import_roster(db,[roster(sid='999901')])
    assert s.students(db,'A')[0]['name']=='新名字'
    assert len(s.students(db,'A'))==1

def test_assignment_swap(db):
    s.import_roster(db,[roster(),roster(sid='90000102',name='测试乙',pocket='02')])
    s.import_roster(db,[roster(pocket='02'),roster(sid='90000102',name='测试乙',pocket='01')])
    assert s.students(db,'A')[0]['sid']=='90000102'

def test_save_dedup_correction_export_snapshot(db):
    s.import_roster(db,[roster(),roster('B')])
    first=s.save_session(db,'A','学期','2026-09-09','第1节','p.jpg',{},[record()])
    again=s.save_session(db,'A','学期','2026-09-09','第1节','p.jpg',{},[record()])
    assert first==again and len(s.details(db))==1
    s.save_session(db,'B','学期','2026-09-09','第1节','p.jpg',{},[record('已交')])
    assert s.summary(db,s.details(db,'A'),'A')[0]['未交次数']==1
    s.import_roster(db,[roster(name='改名',pocket='02')])
    assert s.details(db,'A')[0]['name']=='测试甲' and s.details(db,'A')[0]['pocket']=='01'
    s.correct_record(db,first,'90000101','已交','补交')
    detail=s.details(db,'A'); total=s.summary(db,detail,'A')
    assert total[0]['未交次数']==0 and total[0]['已交次数']==1
    wb=load_workbook(io.BytesIO(s.export_excel(total,detail,'测试')))
    assert wb['明细']['E3'].value=='01' and wb['明细']['I3'].value=='补交'
    assert wb['汇总']['D3'].value==1 and wb['汇总']['E3'].value==0
    assert s.details(db,start='2026-09-10')==[]
    assert db.execute('SELECT COUNT(*) FROM audit').fetchone()[0]==4

def test_review_excluded_and_invalid_save(db):
    s.import_roster(db,[roster()]); s.save_session(db,'A','T','2026-09-09','1','p',{},[record('待复核')])
    t=s.summary(db,s.details(db),'A')[0]; assert t['未交次数']==0 and t['待复核次数']==1
    with pytest.raises(ValueError): s.save_session(db,'A','T','2026-09-09','2','p',{},[record('bad')])
    assert len(s.details(db))==1

def test_excel_multiclass_and_optional_pocket():
    wb=Workbook(); ws=wb.active; ws.append(['学号/工号','姓名','班级','袋号'])
    ws.append(['00101','甲','A',None]); ws.append(['00299','乙','B','02'])
    buffer=io.BytesIO(); wb.save(buffer)
    rows=s.read_roster(buffer.getvalue()); assert rows[0]['sid']=='00101' and rows[1]['pocket']=='02'
    ws.append(['00101','丙','A','03']); buffer=io.BytesIO(); wb.save(buffer)
    with pytest.raises(ValueError): s.read_roster(buffer.getvalue())

def test_formula_text_is_not_executable():
    totals=[{'班级':'=1+1','学号':'001','姓名':'+危险','已交次数':0}]
    wb=load_workbook(io.BytesIO(s.export_excel(totals,[],'过滤')))
    assert wb['汇总']['A3'].data_type=='s'

def test_vision_interface_and_geometry():
    im=read_image(ROOT/'samples/sample_1.jpg'); result,warp=recognize(im)
    assert len(result['pockets'])==54 and warp.shape==(1080,600,3)
    assert result['pockets'][7]['prediction']=='empty'
    assert result['pockets'][0]['prediction']=='phone'
    with pytest.raises(ValueError): rectify(im,[[0,0]]*4)
    with pytest.raises(ValueError): recognize(im,result['corners'],{'rows':[1]*9,'cols':list(range(0,601,100))})

def test_official_demo_separate(tmp_path):
    a=s.connect(tmp_path/'demo/db'); b=s.connect(tmp_path/'official/db')
    s.save_session(a,'A','T','2026-09-09','1','p',{},[record()])
    assert s.details(b)==[]; a.close(); b.close()
